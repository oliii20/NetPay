from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("figure1_netting_reference_style.pptx")


def rgb(value: str) -> RGBColor:
    value = value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


COLORS = {
    "black": rgb("000000"),
    "ink": rgb("243447"),
    "muted": rgb("667085"),
    "blue": rgb("0786C9"),
    "green": rgb("3F6F2A"),
    "red": rgb("D7191C"),
    "teal": rgb("13A89E"),
    "orange": rgb("D17A35"),
    "gold": rgb("CBA636"),
    "purple": rgb("4F6FB8"),
    "light": rgb("F7FAFC"),
    "soft_blue": rgb("EAF6FF"),
    "soft_green": rgb("EEF7E8"),
    "soft_red": rgb("FFF0F0"),
    "soft_teal": rgb("E8FAF7"),
    "gray_line": rgb("D4D9E2"),
}


def set_text(frame, value, size=18, color="ink", bold=False, italic=False, align="center"):
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = alignment
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = COLORS[color]


def text(slide, value, x, y, w, h, size=18, color="ink", bold=False, italic=False, align="left"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box.text_frame, value, size=size, color=color, bold=bold, italic=italic, align=align)
    return box


def box(slide, x, y, w, h, fill="light", line="black", width=2.1, radius=0.04):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[fill]
    shp.line.color.rgb = COLORS[line]
    shp.line.width = Pt(width)
    shp.adjustments[0] = radius
    return shp


def label_box(slide, value, x, y, w, h, fill="light", line="black", size=16, color="ink", bold=True):
    shp = box(slide, x, y, w, h, fill=fill, line=line, width=1.3, radius=0.08)
    set_text(shp.text_frame, value, size=size, color=color, bold=bold)
    return shp


def line(slide, x1, y1, x2, y2, color="black", width=2.0, dashed=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = COLORS[color]
    conn.line.width = Pt(width)
    if dashed:
        conn.line.dash_style = 4
    return conn


def arrow_head(slide, x, y, color="black", rotation=90, size=0.16):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(x - size / 2),
        Inches(y - size / 2),
        Inches(size),
        Inches(size),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[color]
    shp.line.color.rgb = COLORS[color]
    shp.rotation = rotation
    return shp


def arrow(slide, x1, y1, x2, y2, color="black", width=2.0, dashed=False, rotation=90):
    line(slide, x1, y1, x2, y2, color=color, width=width, dashed=dashed)
    arrow_head(slide, x2, y2, color=color, rotation=rotation)


def person(slide, x, y, color="blue", scale=1.0, name=None):
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.10 * scale), Inches(y), Inches(0.28 * scale), Inches(0.28 * scale))
    head.fill.solid()
    head.fill.fore_color.rgb = COLORS[color]
    head.line.color.rgb = COLORS[color]
    body = slide.shapes.add_shape(MSO_SHAPE.TRAPEZOID, Inches(x), Inches(y + 0.28 * scale), Inches(0.48 * scale), Inches(0.42 * scale))
    body.fill.solid()
    body.fill.fore_color.rgb = COLORS[color]
    body.line.color.rgb = COLORS[color]
    body.rotation = 180
    if name:
        text(slide, name, x - 0.08 * scale, y + 0.72 * scale, 0.65 * scale, 0.20 * scale, size=8.8 * scale, align="center")


def receipt(slide, x, y, scale=1.0):
    paper = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.31 * scale), Inches(0.42 * scale))
    paper.fill.solid()
    paper.fill.fore_color.rgb = rgb("FFFFFF")
    paper.line.color.rgb = COLORS["black"]
    paper.line.width = Pt(1.0)
    for i in range(3):
        line(slide, x + 0.05 * scale, y + (0.10 + i * 0.09) * scale, x + 0.25 * scale, y + (0.10 + i * 0.09) * scale, color="black", width=0.7)
    coin = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.24 * scale), Inches(y + 0.16 * scale), Inches(0.18 * scale), Inches(0.18 * scale))
    coin.fill.solid()
    coin.fill.fore_color.rgb = rgb("FFFFFF")
    coin.line.color.rgb = COLORS["black"]
    coin.line.width = Pt(1.0)


def bulb(slide, x, y, scale=1.0):
    lamp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.24 * scale), Inches(0.24 * scale))
    lamp.fill.solid()
    lamp.fill.fore_color.rgb = rgb("FFD84D")
    lamp.line.color.rgb = rgb("F2B705")
    base = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.08 * scale), Inches(y + 0.22 * scale), Inches(0.08 * scale), Inches(0.08 * scale))
    base.fill.solid()
    base.fill.fore_color.rgb = rgb("F2B705")
    base.line.color.rgb = rgb("F2B705")


def pbft_badge(slide, x, y, label, color="black"):
    circ = slide.shapes.add_shape(MSO_SHAPE.CIRCULAR_ARROW, Inches(x), Inches(y), Inches(0.52), Inches(0.42))
    circ.fill.solid()
    circ.fill.fore_color.rgb = COLORS[color]
    circ.line.color.rgb = COLORS[color]
    text(slide, label, x - 0.17, y - 0.26, 0.88, 0.22, size=11, bold=True, align="center")


def shard_titles(slide, x, y, w):
    text(slide, "Shard S1", x + 0.25, y + 0.12, w / 2 - 0.35, 0.28, size=19, color="blue", bold=False, align="center")
    text(slide, "Shard S2", x + w / 2 + 0.1, y + 0.12, w / 2 - 0.35, 0.28, size=19, color="green", bold=False, align="center")
    line(slide, x + w / 2, y + 0.28, x + w / 2, y + 3.25, color="gray_line", width=1.2, dashed=True)


def stage_panel(slide, letter, caption, x, y, w, h):
    box(slide, x, y, w, h, fill="light", line="black", width=2.3, radius=0.035)
    label_box(slide, letter, x + 0.12, y + 0.12, 0.38, 0.30, fill="black", line="black", size=13, color="light")
    text(slide, caption, x + 0.58, y + h + 0.08, w - 0.65, 0.28, size=17, color="ink", align="center")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")

    text(slide, "Figure 1. Bilateral Netting Mechanism", 0.45, 0.18, 7.0, 0.38, size=25, bold=True)
    text(slide, "Full-offset case: reciprocal CTXs become compact settlement txs.", 0.48, 0.58, 6.45, 0.25, size=15, color="muted")

    # Legend.
    box(slide, 7.85, 0.25, 4.85, 0.58, fill="light", line="black", width=1.6, radius=0.02)
    for i, c in enumerate(["black", "orange", "gold", "purple", "teal", "green"]):
        person(slide, 8.04 + i * 0.18, 0.35, color=c, scale=0.44)
    text(slide, "Accounts", 9.15, 0.38, 0.96, 0.20, size=13, align="left")
    receipt(slide, 10.05, 0.32, scale=0.68)
    text(slide, "Intra/Cross-shard\ntransaction (ITX/CTX)", 10.47, 0.30, 2.00, 0.36, size=12.2, align="left")

    # Three panels in the same visual family as the reference.
    panels = [
        (0.45, 1.28, 3.70, 4.75, "(a)", "CTX first half"),
        (4.55, 1.28, 3.70, 4.75, "(b)", "Traditional CTX second half"),
        (8.65, 1.28, 4.20, 4.75, "(c)", "Netting CTX second half"),
    ]
    for x, y, w, h, letter, cap in panels:
        stage_panel(slide, letter, cap, x, y, w, h)

    # Panel a: first-half CTXs.
    x, y, w, h = panels[0][0], panels[0][1], panels[0][2], panels[0][3]
    shard_titles(slide, x + 0.12, y + 0.20, w - 0.24)
    pbft_badge(slide, x + 0.84, y + 0.78, "1. PBFT")
    pbft_badge(slide, x + 2.55, y + 0.78, "1. PBFT")
    person(slide, x + 0.70, y + 1.52, "blue", 1.0, "S1")
    person(slide, x + 2.63, y + 1.52, "green", 1.0, "S2")
    receipt(slide, x + 1.22, y + 1.58, 0.74)
    receipt(slide, x + 2.24, y + 2.43, 0.74)
    arrow(slide, x + 1.62, y + 1.90, x + 2.50, y + 1.90, color="black", width=2.0, dashed=True, rotation=90)
    arrow(slide, x + 2.16, y + 2.72, x + 1.28, y + 2.72, color="black", width=2.0, dashed=True, rotation=270)
    text(slide, "a1, b1, c1", x + 1.58, y + 1.56, 0.95, 0.20, size=14, italic=True, align="center")
    text(slide, "d1, e1, f1", x + 1.35, y + 2.88, 0.95, 0.20, size=14, italic=True, align="center")
    label_box(slide, "n bidirectional\ncross-shard txs", x + 0.72, y + 3.60, 2.25, 0.55, fill="soft_blue", line="black", size=14, color="ink", bold=True)

    # Panel b: traditional second-half CTXs.
    x, y, w, h = panels[1][0], panels[1][1], panels[1][2], panels[1][3]
    shard_titles(slide, x + 0.12, y + 0.20, w - 0.24)
    pbft_badge(slide, x + 0.84, y + 0.78, "2. PBFT")
    pbft_badge(slide, x + 2.55, y + 0.78, "2. PBFT")
    person(slide, x + 0.72, y + 1.50, "blue", 1.0, "S1")
    person(slide, x + 2.62, y + 1.50, "green", 1.0, "S2")
    receipt(slide, x + 2.24, y + 1.58, 0.74)
    receipt(slide, x + 1.22, y + 2.43, 0.74)
    arrow(slide, x + 2.15, y + 1.90, x + 1.30, y + 1.90, color="red", width=2.2, dashed=False, rotation=270)
    arrow(slide, x + 1.60, y + 2.72, x + 2.47, y + 2.72, color="red", width=2.2, dashed=False, rotation=90)
    text(slide, "a2, b2, c2", x + 1.45, y + 1.56, 0.90, 0.20, size=14, italic=True, align="center")
    text(slide, "d2, e2, f2", x + 1.48, y + 2.88, 0.90, 0.20, size=14, italic=True, align="center")
    label_box(slide, "Still process\nn second-half CTXs", x + 0.55, y + 3.60, 2.65, 0.55, fill="soft_red", line="red", size=14, color="red", bold=True)

    # Panel c: netting settlement.
    x, y, w, h = panels[2][0], panels[2][1], panels[2][2], panels[2][3]
    shard_titles(slide, x + 0.12, y + 0.20, w - 0.24)
    pbft_badge(slide, x + 0.94, y + 0.78, "2. PBFT")
    pbft_badge(slide, x + 2.90, y + 0.78, "2. PBFT")
    person(slide, x + 0.82, y + 1.38, "blue", 1.0)
    person(slide, x + 2.92, y + 1.38, "green", 1.0)
    solver = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 1.88), Inches(y + 2.58), Inches(0.55), Inches(0.55))
    solver.fill.solid()
    solver.fill.fore_color.rgb = rgb("FFFFFF")
    solver.line.color.rgb = COLORS["black"]
    solver.line.width = Pt(1.8)
    gear = slide.shapes.add_shape(MSO_SHAPE.FLOWCHART_PREPARATION, Inches(x + 1.58), Inches(y + 3.10), Inches(1.14), Inches(0.42))
    gear.fill.solid()
    gear.fill.fore_color.rgb = COLORS["soft_teal"]
    gear.line.color.rgb = COLORS["black"]
    gear.line.width = Pt(1.6)
    set_text(gear.text_frame, "Solver", size=12.5, color="ink", bold=True)
    receipt(slide, x + 1.08, y + 2.53, 0.68)
    receipt(slide, x + 2.78, y + 2.53, 0.68)
    arrow(slide, x + 2.05, y + 2.58, x + 1.25, y + 2.03, color="green", width=2.4, dashed=False, rotation=310)
    arrow(slide, x + 2.25, y + 2.58, x + 3.03, y + 2.03, color="green", width=2.4, dashed=False, rotation=50)
    text(slide, "settlement\ntx", x + 0.48, y + 2.14, 1.00, 0.38, size=11.8, italic=True, align="center")
    text(slide, "settlement\ntx", x + 2.82, y + 2.14, 1.00, 0.38, size=11.8, italic=True, align="center")
    label_box(slide, "Full offset:\nno CTX residual", x + 0.42, y + 3.72, 2.10, 0.58, fill="soft_teal", line="green", size=14.0, color="green", bold=True)

    # Effect callout.
    box(slide, 11.42, 4.48, 1.24, 0.78, fill="light", line="black", width=1.8, radius=0.12)
    bulb(slide, 11.51, 4.68, 0.66)
    text(slide, "Low CTX\nload", 11.67, 4.72, 0.80, 0.30, size=12.0, color="red", bold=True, align="center")

    prs.save(OUT)


if __name__ == "__main__":
    build()
