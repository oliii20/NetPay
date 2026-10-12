from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("figure1_bilateral_netting_best_case.pptx")


def rgb(hex_value: str) -> RGBColor:
    hex_value = hex_value.strip("#")
    return RGBColor(
        int(hex_value[0:2], 16),
        int(hex_value[2:4], 16),
        int(hex_value[4:6], 16),
    )


COLORS = {
    "ink": rgb("1F2937"),
    "muted": rgb("6B7280"),
    "light": rgb("F8FAFC"),
    "panel": rgb("FFFFFF"),
    "line": rgb("CBD5E1"),
    "shard": rgb("E0F2FE"),
    "shard_line": rgb("0369A1"),
    "trad": rgb("DC2626"),
    "trad_soft": rgb("FEE2E2"),
    "net": rgb("059669"),
    "net_soft": rgb("D1FAE5"),
    "solver": rgb("EEF2FF"),
    "solver_line": rgb("4F46E5"),
    "amber": rgb("F59E0B"),
    "amber_soft": rgb("FEF3C7"),
}


def add_text(slide, text, x, y, w, h, size=10, color="ink", bold=False, align="left"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(frame, text, size=size, color=color, bold=bold, align=align)
    return box


def set_text(frame, text, size=10, color="ink", bold=False, align="center"):
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    align_value = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(lines):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align_value
        p.space_after = Pt(0)
        p.space_before = Pt(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = COLORS[color]


def add_label(slide, text, x, y, w, h, fill, line=None, size=9, color="ink", bold=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[fill]
    shape.line.color.rgb = COLORS[line or fill]
    shape.line.width = Pt(0.9)
    shape.adjustments[0] = 0.18
    tf = shape.text_frame
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.03)
    tf.margin_bottom = Inches(0.03)
    set_text(tf, text, size=size, color=color, bold=bold)
    return shape


def add_rect(slide, text, x, y, w, h, fill, line, size=12, bold=True):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[fill]
    shape.line.color.rgb = COLORS[line]
    shape.line.width = Pt(1.4)
    shape.adjustments[0] = 0.1
    tf = shape.text_frame
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    set_text(tf, text, size=size, color="ink", bold=bold)
    return shape


def add_panel(slide, label, title, x, y, w, h):
    panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = COLORS["panel"]
    panel.line.color.rgb = COLORS["line"]
    panel.line.width = Pt(0.9)
    panel.adjustments[0] = 0.04
    add_label(slide, label, x + 0.15, y + 0.14, 0.36, 0.28, "ink", "ink", size=8, color="panel", bold=True)
    add_text(slide, title, x + 0.57, y + 0.12, w - 0.75, 0.32, size=10.5, bold=True)
    return panel


def add_arrow(slide, x1, y1, x2, y2, color, width=2.0, begin=False, end=True, dash=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    line = conn.line
    line.color.rgb = COLORS[color]
    line.width = Pt(width)
    if end:
        line.end_arrowhead = True
    if begin:
        line.begin_arrowhead = True
    if dash:
        line.dash_style = 4
    return conn


def add_metric(slide, title, value, x, y, w, color):
    add_label(slide, title, x, y, w, 0.28, "light", "line", size=7.4, color="muted", bold=True)
    add_label(slide, value, x, y + 0.31, w, 0.42, color, color, size=11.5, color="ink", bold=True)


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = COLORS["light"]

    add_text(
        slide,
        "Figure 1. Bilateral netting removes reciprocal cross-shard workload in the best case",
        0.5,
        0.25,
        12.3,
        0.42,
        size=18,
        bold=True,
    )
    add_text(
        slide,
        "Balanced opposite flows are aggregated before settlement, so only the net residual enters the cross-shard confirmation path.",
        0.52,
        0.66,
        7.35,
        0.28,
        size=8.8,
        color="muted",
    )

    add_panel(slide, "a", "Traditional cross-shard processing", 0.55, 1.12, 5.95, 4.5)
    add_panel(slide, "b", "Proposed bilateral netting, full-offset case", 6.82, 1.12, 5.95, 4.5)

    # Traditional panel.
    add_rect(slide, "Shard S1\nPBFT group", 1.05, 2.18, 1.25, 0.9, "shard", "shard_line", size=11)
    add_rect(slide, "Shard S2\nPBFT group", 4.72, 2.18, 1.25, 0.9, "shard", "shard_line", size=11)
    add_label(slide, "local consensus", 1.08, 1.83, 1.18, 0.28, "light", "line", size=7, color="muted")
    add_label(slide, "local consensus", 4.76, 1.83, 1.18, 0.28, "light", "line", size=7, color="muted")

    for dy in [0.0, 0.17, 0.34, 0.51]:
        add_arrow(slide, 2.34, 2.25 + dy, 4.65, 2.25 + dy, "trad", width=1.8)
    for dy in [0.0, 0.17, 0.34, 0.51]:
        add_arrow(slide, 4.65, 3.13 + dy, 2.34, 3.13 + dy, "trad", width=1.8)
    add_label(slide, "100 cross-shard tx\nS1 to S2", 2.52, 1.83, 1.28, 0.52, "trad_soft", "trad", size=8.2, color="trad", bold=True)
    add_label(slide, "100 cross-shard tx\nS2 to S1", 3.18, 3.72, 1.28, 0.52, "trad_soft", "trad", size=8.2, color="trad", bold=True)
    add_label(slide, "Both directions enter relay / broker / receipt handling", 1.18, 4.38, 4.65, 0.38, "amber_soft", "amber", size=8.7, color="ink", bold=True)
    add_text(slide, "Cross-shard workload = 200 tx-equivalents", 1.25, 4.88, 4.45, 0.26, size=9, color="muted", align="center")

    # Netting panel.
    add_rect(slide, "Shard S1\nPBFT group", 7.15, 2.18, 1.25, 0.9, "shard", "shard_line", size=11)
    add_rect(slide, "Shard S2\nPBFT group", 11.18, 2.18, 1.25, 0.9, "shard", "shard_line", size=11)
    solver = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(8.93), Inches(2.05), Inches(1.75), Inches(1.15))
    solver.fill.solid()
    solver.fill.fore_color.rgb = COLORS["solver"]
    solver.line.color.rgb = COLORS["solver_line"]
    solver.line.width = Pt(1.4)
    solver.text_frame.clear()
    solver.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    set_text(solver.text_frame, "Netting\nsolver", size=11, color="ink", bold=True)

    add_arrow(slide, 8.43, 2.42, 8.93, 2.42, "trad", width=2.0)
    add_arrow(slide, 11.14, 2.82, 10.67, 2.82, "trad", width=2.0)
    add_label(slide, "Gross S1 to S2 = 100", 8.22, 1.64, 1.7, 0.32, "trad_soft", "trad", size=7.6, color="trad", bold=True)
    add_label(slide, "Gross S2 to S1 = 100", 9.85, 3.28, 1.7, 0.32, "trad_soft", "trad", size=7.6, color="trad", bold=True)

    add_label(slide, "cancel reciprocal obligations", 8.62, 3.64, 2.38, 0.35, "solver", "solver_line", size=8.2, color="solver_line", bold=True)
    add_label(slide, "Net residual = 0", 8.72, 4.13, 2.15, 0.46, "net_soft", "net", size=12, color="net", bold=True)
    add_arrow(slide, 9.75, 4.6, 9.75, 4.98, "net", width=2.2)
    add_label(slide, "No ordinary cross-shard settlement tx", 8.37, 4.97, 2.78, 0.38, "net_soft", "net", size=8.8, color="net", bold=True)
    add_arrow(slide, 9.05, 3.18, 8.3, 3.62, "solver_line", width=1.2, dash=True)
    add_arrow(slide, 10.55, 3.18, 11.3, 3.62, "solver_line", width=1.2, dash=True)
    add_text(slide, "compact clearing certificate", 8.75, 3.17, 2.0, 0.24, size=7.2, color="muted", align="center")

    # Outcome strip.
    add_panel(slide, "c", "Mechanism-level effect", 0.55, 5.87, 12.22, 1.04)
    add_metric(slide, "cross-shard settlement", "200 -> 0", 1.05, 6.28, 1.65, "net_soft")
    add_metric(slide, "PBFT/block pressure", "lower", 3.2, 6.28, 1.65, "net_soft")
    add_metric(slide, "relay/broker messages", "lower", 5.35, 6.28, 1.65, "net_soft")
    add_metric(slide, "txpool tail latency", "shorter", 7.5, 6.28, 1.65, "net_soft")
    add_text(
        slide,
        "Best case shown: perfectly balanced bilateral flows. Netting benefit decreases as traffic becomes one-directional.",
        9.75,
        6.19,
        2.45,
        0.55,
        size=7.5,
        color="muted",
    )

    # Small equation callout.
    eq = add_label(slide, "min(100, 100) is offset; |100 - 100| = 0 remains", 8.25, 0.59, 4.25, 0.32, "net_soft", "net", size=8.0, color="net", bold=True)
    eq.line.width = Pt(0.8)

    prs.save(OUT)


if __name__ == "__main__":
    build()
