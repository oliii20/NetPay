from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("figure1_netting_sketch_replica.pptx")


def rgb(value: str) -> RGBColor:
    value = value.strip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


COLORS = {
    "ink": rgb("111827"),
    "muted": rgb("6B7280"),
    "line": rgb("111827"),
    "soft_line": rgb("CBD5E1"),
    "shard": rgb("EAF6FF"),
    "shard_line": rgb("2563EB"),
    "trad": rgb("DC2626"),
    "trad_soft": rgb("FEE2E2"),
    "net": rgb("059669"),
    "net_soft": rgb("D1FAE5"),
    "note": rgb("F8FAFC"),
}


def set_text(frame, text, size=10, color="ink", bold=False, italic=False, align="center"):
    frame.clear()
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    align_value = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = align_value
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = COLORS[color]


def text(slide, value, x, y, w, h, size=10, color="ink", bold=False, italic=False, align="left"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box.text_frame, value, size=size, color=color, bold=bold, italic=italic, align=align)
    return box


def rounded_box(slide, value, x, y, w, h, fill, line, size=13, bold=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[fill]
    shp.line.color.rgb = COLORS[line]
    shp.line.width = Pt(1.4)
    shp.adjustments[0] = 0.08
    set_text(shp.text_frame, value, size=size, color="ink", bold=bold)
    return shp


def note_box(slide, value, x, y, w, h, size=9.2, color="ink", bold=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS["note"]
    shp.line.color.rgb = COLORS["soft_line"]
    shp.line.width = Pt(0.8)
    shp.adjustments[0] = 0.07
    set_text(shp.text_frame, value, size=size, color=color, bold=bold)
    return shp


def line(slide, x1, y1, x2, y2, color="line", width=1.5, end=True, begin=False, dash=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = COLORS[color]
    conn.line.width = Pt(width)
    if end:
        conn.line.end_arrowhead = True
    if begin:
        conn.line.begin_arrowhead = True
    if dash:
        conn.line.dash_style = 4
    return conn


def arrow_head(slide, x, y, color="line", rotation=90, size=0.13):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(x - size / 2),
        Inches(y - size / 2),
        Inches(size),
        Inches(size),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[color]
    shp.line.color.rgb = COLORS[color]
    shp.rotation = rotation
    return shp


def h_arrow(slide, x1, y, x2, color="line", width=1.5):
    line(slide, x1, y, x2, y, color=color, width=width, end=False)
    arrow_head(slide, x2, y, color=color, rotation=90 if x2 > x1 else 270)


def d_arrow(slide, x1, y1, x2, y2, color="line", width=1.5, rotation=45):
    line(slide, x1, y1, x2, y2, color=color, width=width, end=False)
    arrow_head(slide, x2, y2, color=color, rotation=rotation)


def pbft_loop(slide, x, y, label):
    arc = slide.shapes.add_shape(MSO_SHAPE.CURVED_DOWN_ARROW, Inches(x), Inches(y), Inches(0.62), Inches(0.38))
    arc.fill.solid()
    arc.fill.fore_color.rgb = COLORS["line"]
    arc.line.color.rgb = COLORS["line"]
    text(slide, label, x - 0.08, y - 0.28, 0.85, 0.22, size=8.7, bold=True, align="center")
    return arc


def section_rule(slide, y):
    line(slide, 0.6, y, 6.9, y, color="soft_line", width=0.8, end=False)


def shard_pair(slide, y, pbft_label, x_left=2.2, x_right=5.0, box_w=0.95, box_h=0.7):
    s1 = rounded_box(slide, "S1", x_left, y, box_w, box_h, "shard", "shard_line", size=14)
    s2 = rounded_box(slide, "S2", x_right, y, box_w, box_h, "shard", "shard_line", size=14)
    pbft_loop(slide, x_left + 0.18, y - 0.46, pbft_label)
    pbft_loop(slide, x_right + 0.18, y - 0.46, pbft_label)
    return s1, s2


def build():
    prs = Presentation()
    prs.slide_width = Inches(7.5)
    prs.slide_height = Inches(10.0)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")

    text(slide, "Figure 1. Bilateral netting mechanism under full offset", 0.55, 0.28, 6.4, 0.34, size=16, bold=True)
    text(slide, "First-half CTXs, traditional second-half CTXs, and netting-based settlement.", 0.57, 0.65, 6.3, 0.26, size=7.8, color="muted")
    section_rule(slide, 1.05)

    # Top: first half.
    text(slide, "CTX first half", 0.65, 1.27, 1.35, 0.28, size=11.5, bold=True)
    shard_pair(slide, 1.95, "1. PBFT")
    h_arrow(slide, 3.20, 2.13, 4.84, color="line", width=1.2)
    h_arrow(slide, 4.94, 2.42, 3.30, color="line", width=1.2)
    text(slide, "a1, b1, c1", 3.37, 1.86, 1.28, 0.22, size=8.7, italic=True, align="center")
    text(slide, "d1, e1, f1", 3.37, 2.45, 1.28, 0.22, size=8.7, italic=True, align="center")
    note_box(slide, "There are n bidirectional\ncross-shard transactions\nbetween S1 and S2.", 5.95, 1.85, 1.05, 0.86, size=8.0)
    section_rule(slide, 3.05)

    # Middle: traditional second half.
    text(slide, "CTX second half", 0.65, 3.25, 1.55, 0.28, size=11.5, bold=True)
    text(slide, "Traditional:", 0.9, 3.7, 1.25, 0.35, size=12.0, italic=True, bold=True)
    shard_pair(slide, 4.15, "2. PBFT", x_left=2.05, x_right=4.72)
    h_arrow(slide, 4.60, 4.32, 3.10, color="trad", width=1.4)
    h_arrow(slide, 3.10, 4.60, 4.60, color="trad", width=1.4)
    text(slide, "a2, b2, c2", 3.15, 4.07, 1.1, 0.22, size=8.7, italic=True, align="center")
    text(slide, "d2, e2, f2", 3.15, 4.63, 1.1, 0.22, size=8.7, italic=True, align="center")
    note_box(slide, "Traditional protocols still process\nn second-half CTXs.", 1.58, 5.26, 4.55, 0.46, size=9.3)
    section_rule(slide, 6.05)

    # Bottom: netting.
    text(slide, "Netting:", 0.9, 6.35, 1.25, 0.34, size=12.0, italic=True, bold=True)
    rounded_box(slide, "S1", 2.03, 6.98, 0.82, 0.65, "shard", "shard_line", size=14)
    rounded_box(slide, "S2", 4.82, 6.98, 0.82, 0.65, "shard", "shard_line", size=14)
    pbft_loop(slide, 2.08, 6.52, "2. PBFT")
    pbft_loop(slide, 4.87, 6.52, "2. PBFT")

    solver = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.38), Inches(8.15), Inches(0.46), Inches(0.46))
    solver.fill.solid()
    solver.fill.fore_color.rgb = rgb("FFFFFF")
    solver.line.color.rgb = COLORS["net"]
    solver.line.width = Pt(1.5)
    body = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(3.33), Inches(8.56), Inches(0.56), Inches(0.5))
    body.fill.solid()
    body.fill.fore_color.rgb = rgb("FFFFFF")
    body.line.color.rgb = COLORS["net"]
    body.line.width = Pt(1.5)
    text(slide, "Solver", 3.1, 9.06, 1.05, 0.25, size=10.2, italic=True, bold=True, align="center")

    d_arrow(slide, 3.56, 8.14, 2.52, 7.72, color="net", width=1.8, rotation=315)
    d_arrow(slide, 3.72, 8.14, 5.13, 7.72, color="net", width=1.8, rotation=45)
    text(slide, "x settlement tx", 1.18, 7.72, 1.55, 0.25, size=9.0, italic=True, align="center")
    text(slide, "settlement tx", 4.82, 7.72, 1.45, 0.25, size=9.0, italic=True, align="center")
    note_box(slide, "Under full bilateral offset,\nNetting only issues compact\nsettlement transactions.", 2.05, 9.42, 3.55, 0.52, size=8.8, color="net", bold=True)

    prs.save(OUT)


if __name__ == "__main__":
    build()
