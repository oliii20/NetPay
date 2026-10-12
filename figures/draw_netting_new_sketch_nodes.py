from pathlib import Path
from math import cos, sin, pi

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("figure1_netting_new_sketch_nodes.pptx")


def rgb(value: str) -> RGBColor:
    value = value.strip("#")
    return RGBColor(int(value[:2], 16), int(value[2:4], 16), int(value[4:], 16))


COLORS = {
    "ink": rgb("111827"),
    "muted": rgb("667085"),
    "black": rgb("000000"),
    "blue": rgb("2F62D6"),
    "red": rgb("F01818"),
    "green": rgb("078B5B"),
    "green_soft": rgb("E9F8F0"),
    "red_soft": rgb("FFF0F0"),
    "blue_soft": rgb("EFF6FF"),
    "panel": rgb("F8FAFC"),
    "line": rgb("D0D5DD"),
    "yellow": rgb("FFD84D"),
}


def set_text(frame, text, size=16, color="ink", bold=False, italic=False, align="center"):
    frame.clear()
    frame.margin_left = Inches(0.04)
    frame.margin_right = Inches(0.04)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    for i, line in enumerate(text.split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.alignment = alignment
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = COLORS[color]


def text(slide, value, x, y, w, h, **kwargs):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    set_text(box.text_frame, value, **kwargs)
    return box


def connector(slide, x1, y1, x2, y2, color="black", width=1.4, dashed=False):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT,
        Inches(x1),
        Inches(y1),
        Inches(x2),
        Inches(y2),
    )
    conn.line.color.rgb = COLORS[color]
    conn.line.width = Pt(width)
    if dashed:
        conn.line.dash_style = 4
    return conn


def arrow_head(slide, x, y, color="black", rotation=90, size=0.15):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ISOSCELES_TRIANGLE,
        Inches(x - size / 2),
        Inches(y - size / 2),
        Inches(size),
        Inches(size),
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[color]
    shp.line.color.rgb = COLORS[color]
    shp.rotation = rotation
    return shp


def arrow(slide, x1, y1, x2, y2, color="black", width=1.8, dashed=False, rotation=90):
    connector(slide, x1, y1, x2, y2, color=color, width=width, dashed=dashed)
    arrow_head(slide, x2, y2, color=color, rotation=rotation)


def rounded_rect(slide, x, y, w, h, fill="panel", line="black", width=1.8, radius=0.05):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[fill]
    shp.line.color.rgb = COLORS[line]
    shp.line.width = Pt(width)
    shp.adjustments[0] = radius
    return shp


def section_label(slide, label, x, y, color="ink"):
    text(slide, label, x, y, 2.85, 0.48, size=16.5, color=color, bold=True, align="left")


def node(slide, x, y, color, size=0.16):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x - size / 2), Inches(y - size / 2), Inches(size), Inches(size))
    shp.fill.solid()
    shp.fill.fore_color.rgb = COLORS[color]
    shp.line.color.rgb = COLORS[color]
    shp.line.width = Pt(0.8)
    return shp


def draw_shard(slide, cx, cy, r, color, label):
    outer = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2 * r), Inches(2 * r))
    outer.fill.background()
    outer.line.color.rgb = COLORS[color]
    outer.line.width = Pt(1.8)
    outer.line.dash_style = 3

    positions = [
        (-0.45, -0.10),
        (-0.12, -0.43),
        (0.28, -0.12),
        (-0.15, 0.18),
        (0.18, 0.36),
    ]
    pts = [(cx + dx * r, cy + dy * r) for dx, dy in positions]
    for i, j in [(0, 1), (1, 2), (2, 3), (3, 0), (0, 2), (1, 3), (3, 4), (2, 4)]:
        connector(slide, pts[i][0], pts[i][1], pts[j][0], pts[j][1], color=color, width=1.1)
    for px, py in pts:
        node(slide, px, py, color, size=0.17)

    text(slide, label, cx - 0.55, cy + r + 0.08, 1.1, 0.25, size=15, color="ink", bold=True, align="center")


def solver_icon(slide, x, y):
    head = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.25), Inches(0.25))
    head.fill.solid()
    head.fill.fore_color.rgb = COLORS["green_soft"]
    head.line.color.rgb = COLORS["green"]
    head.line.width = Pt(1.4)
    body = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(x - 0.05), Inches(y + 0.23), Inches(0.35), Inches(0.36))
    body.fill.solid()
    body.fill.fore_color.rgb = COLORS["green_soft"]
    body.line.color.rgb = COLORS["green"]
    body.line.width = Pt(1.4)
    text(slide, "Solver", x - 0.35, y + 0.60, 0.95, 0.22, size=13, color="green", bold=True, align="center")


def callout(slide, x, y, w, h, title, body, fill="panel", accent="black"):
    rounded_rect(slide, x, y, w, h, fill=fill, line=accent, width=1.6, radius=0.08)
    tri = slide.shapes.add_shape(MSO_SHAPE.RIGHT_TRIANGLE, Inches(x - 0.24), Inches(y + 0.22), Inches(0.34), Inches(0.34))
    tri.fill.solid()
    tri.fill.fore_color.rgb = COLORS[fill]
    tri.line.color.rgb = COLORS[accent]
    tri.line.width = Pt(1.2)
    tri.rotation = 225
    text(slide, title, x + 0.16, y + 0.07, w - 0.32, 0.22, size=14, color=accent, bold=True, align="left")
    text(slide, body, x + 0.16, y + 0.31, w - 0.32, h - 0.34, size=12.4, color="ink", align="left")


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")

    text(slide, "Figure 1. Bilateral Netting Mechanism", 0.45, 0.18, 8.8, 0.36, size=24, bold=True, align="left")
    text(slide, "Shard 1 nodes are blue; Shard 2 nodes are red. Reciprocal source-side CTXs are offset before destination-side execution.", 0.47, 0.58, 9.1, 0.24, size=13.5, color="muted", align="left")

    # Legend inspired by the reference figure.
    rounded_rect(slide, 10.05, 0.18, 2.65, 0.56, fill="panel", line="black", width=1.4, radius=0.02)
    node(slide, 10.35, 0.46, "blue", size=0.15)
    text(slide, "Shard 1 node", 10.50, 0.34, 0.78, 0.22, size=11.5, align="left")
    node(slide, 11.36, 0.46, "red", size=0.15)
    text(slide, "Shard 2 node", 11.51, 0.34, 0.78, 0.22, size=11.5, align="left")

    rows = [
        (1.08, "Source-side CTX:", "source"),
        (3.18, "Destination-side CTX:\nTraditional", "traditional"),
        (5.28, "Netting (proposed):", "netting"),
    ]

    for y, label, _ in rows:
        connector(slide, 0.55, y - 0.10, 12.75, y - 0.10, color="line", width=0.9)
        section_label(slide, label, 0.62, y, color="ink")

    # Row 1: source-side CTXs.
    y = 1.95
    draw_shard(slide, 1.78, y, 0.58, "blue", "Shard 1")
    draw_shard(slide, 4.12, y, 0.58, "red", "Shard 2")
    arrow(slide, 2.38, y - 0.12, 3.48, y - 0.12, color="black", width=1.7, dashed=False, rotation=90)
    arrow(slide, 3.48, y + 0.16, 2.38, y + 0.16, color="black", width=1.7, dashed=False, rotation=270)
    text(slide, "a1, b1, c1, ...", 2.58, y - 0.40, 0.95, 0.22, size=12.5, italic=True, align="center")
    text(slide, "d1, e1, f1, ...", 2.58, y + 0.32, 0.95, 0.22, size=12.5, italic=True, align="center")
    callout(
        slide,
        5.35,
        1.38,
        3.15,
        0.86,
        "Observation",
        "k reciprocal source-side CTXs are generated between Shard 1 and Shard 2.",
        fill="blue_soft",
        accent="blue",
    )

    # Row 2: traditional destination-side CTXs.
    y = 4.05
    draw_shard(slide, 1.78, y, 0.58, "blue", "Shard 1")
    draw_shard(slide, 4.12, y, 0.58, "red", "Shard 2")
    arrow(slide, 3.48, y - 0.12, 2.38, y - 0.12, color="red", width=2.0, dashed=False, rotation=270)
    arrow(slide, 2.38, y + 0.16, 3.48, y + 0.16, color="red", width=2.0, dashed=False, rotation=90)
    text(slide, "a2, b2, c2, ...", 2.55, y - 0.40, 1.00, 0.22, size=12.5, italic=True, align="center")
    text(slide, "d2, e2, f2, ...", 2.55, y + 0.32, 1.00, 0.22, size=12.5, italic=True, align="center")
    callout(
        slide,
        5.35,
        3.48,
        3.25,
        0.86,
        "Traditional",
        "Still processes k destination-side CTXs through PBFT and block packing.",
        fill="red_soft",
        accent="red",
    )

    # Row 3: proposed netting.
    y = 6.15
    draw_shard(slide, 1.55, y, 0.50, "blue", "Shard 1")
    draw_shard(slide, 4.35, y, 0.50, "red", "Shard 2")
    solver_icon(slide, 2.80, 6.42)
    arrow(slide, 2.93, 6.42, 1.78, 6.32, color="green", width=2.0, dashed=False, rotation=280)
    arrow(slide, 3.05, 6.42, 4.18, 6.32, color="green", width=2.0, dashed=False, rotation=80)
    text(slide, "settlement tx", 1.98, 6.34, 1.02, 0.20, size=11.5, color="green", italic=True, align="center")
    text(slide, "settlement tx", 3.20, 6.34, 1.02, 0.20, size=11.5, color="green", italic=True, align="center")
    callout(
        slide,
        5.35,
        5.58,
        3.38,
        0.88,
        "Proposed Netting",
        "Offsets reciprocal CTXs first; only compact settlement CTXs remain (a << k).",
        fill="green_soft",
        accent="green",
    )

    # Main effect box on the right.
    rounded_rect(slide, 9.20, 2.22, 3.32, 2.18, fill="panel", line="black", width=1.8, radius=0.08)
    bulb = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9.48), Inches(2.49), Inches(0.30), Inches(0.30))
    bulb.fill.solid()
    bulb.fill.fore_color.rgb = COLORS["yellow"]
    bulb.line.color.rgb = rgb("E5A400")
    text(slide, "Mechanism insight", 9.86, 2.44, 1.75, 0.28, size=15, bold=True, align="left")
    text(
        slide,
        "Netting offsets reciprocal\nsource-side CTXs before\ndestination-side execution.\n\nOnly a compact settlement set\n(a << k) reaches consensus.",
        9.50,
        2.84,
        2.66,
        1.34,
        size=12.6,
        color="ink",
        align="left",
    )

    prs.save(OUT)


if __name__ == "__main__":
    build()
